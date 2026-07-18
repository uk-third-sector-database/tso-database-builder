"""Social Housing England: registered providers of social housing.

Source page : https://www.gov.uk/government/publications/current-registered-providers-of-social-housing
Product     : "List of registered providers" spreadsheet (updated monthly
              by the Regulator of Social Housing)
Output      : <raw_root>/SocialHousingEngland/registered_providers_YYYYMMDD.csv

The publication page links one .xlsx per month (e.g.
List_of_registered_providers_29_June_2026.xlsx). This module finds the
latest spreadsheet link, reads the providers sheet (the sheet whose
header row starts 'Organisation name'; the workbook also contains a
hidden lookup sheet which is skipped), and writes it to CSV with dates
formatted DD/MM/YYYY - the format handler/socialHousingEng.py parses.
The snapshot date in the output name comes from the spreadsheet's
filename, falling back to today's date.
"""

import argparse
import csv
import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

from .common import DEFAULT_RAW_ROOT, http_get, resolve_outdir

PUBLICATION_PAGE = ('https://www.gov.uk/government/publications/'
                    'current-registered-providers-of-social-housing')


def find_latest_spreadsheet():
    """Return (url, snapshot_date) for the latest providers spreadsheet
    linked from the gov.uk publication page."""
    resp = http_get(PUBLICATION_PAGE, timeout=60)
    links = re.findall(
        r'href="(https://assets\.publishing\.service\.gov\.uk/[^"]+\.xlsx)"',
        resp.text)
    if not links:
        raise RuntimeError(
            f'No .xlsx link found on {PUBLICATION_PAGE}')
    url = links[0]
    snap = _date_from_filename(url) or date.today()
    return url, snap


def _date_from_filename(url):
    """Parse a date like 29_June_2026 (or 29_06_2026) out of the asset
    filename. Returns a date or None."""
    name = url.rsplit('/', 1)[-1]
    m = re.search(r'(\d{1,2})_([A-Za-z]+|\d{2})_(\d{4})', name)
    if not m:
        return None
    d, mon, y = m.groups()
    for fmt in ('%d_%B_%Y', '%d_%b_%Y', '%d_%m_%Y'):
        try:
            return datetime.strptime(f'{d}_{mon}_{y}', fmt).date()
        except ValueError:
            continue
    return None


def _format_cell(value):
    """Spreadsheet cell -> CSV text; dates become DD/MM/YYYY."""
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%d/%m/%Y')
    if isinstance(value, date):
        return value.strftime('%d/%m/%Y')
    return str(value).strip()


def download_providers(outdir=None):
    """Download the latest registered-providers spreadsheet and convert
    the providers sheet to CSV. Returns the path written."""
    folder = resolve_outdir(outdir, 'SocialHousingEngland')
    url, snap = find_latest_spreadsheet()
    print(f'Downloading {url} (snapshot {snap})')
    resp = http_get(url)
    print(f'  got {len(resp.content):,} bytes')

    wb = load_workbook(io.BytesIO(resp.content), read_only=True)
    sheet = None
    for name in wb.sheetnames:
        ws = wb[name]
        first_row = next(ws.iter_rows(values_only=True), None)
        if first_row and str(first_row[0]).strip() == 'Organisation name':
            sheet = ws
            break
    if sheet is None:
        raise RuntimeError(
            f'No sheet with an "Organisation name" header found; sheets '
            f'are {wb.sheetnames} - the layout has changed')

    out = folder / f'registered_providers_{snap:%Y%m%d}.csv'
    n = 0
    with open(out, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        for row in sheet.iter_rows(values_only=True):
            values = [_format_cell(c) for c in row]
            if not any(values):
                continue        # skip blank rows
            writer.writerow(values)
            n += 1
    print(f'  wrote {n - 1} provider rows -> {out}')
    return out


if __name__ == '__main__':
    parser = argparse.ArgumentParser(
        description='Download the monthly "registered providers of social '
                    'housing" list and convert it to CSV')
    parser.add_argument('--outdir', default=None,
                        help=f'raw-data root (default {DEFAULT_RAW_ROOT})')
    args = parser.parse_args()
    download_providers(outdir=args.outdir)
